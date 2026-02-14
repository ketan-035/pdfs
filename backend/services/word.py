import docx

def read_word(path):
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs)

def save_word(text, path):
    d = docx.Document()

    for line in text.split("\n"):
        d.add_paragraph(line)

    d.save(path)
